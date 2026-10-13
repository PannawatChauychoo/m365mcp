"""Read cells from an .xlsx/.xlsm file held in memory.

Excel attachments on mail are not drive items, so the Graph Workbook API
cannot open them; the graph client downloads the bytes and this module parses
them locally with openpyxl in read-only mode. Nothing here touches the network.

The bytes come from arbitrary senders, so the archive is inflated and checked
by zip_guard before openpyxl sees it (size, compression ratio, member count,
and no XML DTDs in UTF-8 or UTF-16), and openpyxl parses with defusedxml. ``read_workbook`` is synchronous and
CPU-bound; callers run it in a worker thread.
"""

from __future__ import annotations

import io
import re
import warnings
from datetime import date, datetime, time, timedelta
from time import monotonic
from typing import Any, Callable, NamedTuple

try:  # pragma: no cover - spreadsheet reading degrades to a reason string
    import openpyxl
    from openpyxl.utils.cell import get_column_letter, range_boundaries
except Exception:  # pragma: no cover
    openpyxl = None  # type: ignore[assignment]

from .models import AttachmentRangeData, WorkbookDefinedNameInfo, WorkbookSheetInfo
from .zip_guard import NotAZipError, ZipGuardError, check_zip_archive

# openpyxl warns about every Excel feature it does not model (data validation
# extensions, slicers, ...); none affect values. A process-wide filter rather
# than warnings.catch_warnings, which is not thread-safe.
warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl(\.|$)")

# Only the OOXML formats openpyxl reads are accepted; legacy .xls and binary
# .xlsb get a reason instead.
SPREADSHEET_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.template",
    "application/vnd.ms-excel.sheet.macroenabled.12",
    "application/vnd.ms-excel.template.macroenabled.12",
}
SPREADSHEET_EXTENSIONS = (".xlsx", ".xlsm", ".xltx", ".xltm")
LEGACY_SPREADSHEET_CONTENT_TYPES = {
    "application/vnd.ms-excel",
    "application/vnd.ms-excel.sheet.binary.macroenabled.12",
}
LEGACY_SPREADSHEET_EXTENSIONS = (".xls", ".xlsb", ".xlt")
DEFAULT_WORKBOOK_MAX_BYTES = 10_000_000
MAX_WORKBOOK_MAX_BYTES = 25_000_000
# Cell values land in the model's context, so one call is bounded well below
# what a sheet can hold; page through a large sheet by range instead.
DEFAULT_WORKBOOK_MAX_CELLS = 20_000
MAX_WORKBOOK_MAX_CELLS = 50_000
# An .xlsx is a zip of XML. Real workbooks inflate to roughly 5-20x their size;
# a zip bomb inflates by thousands. Both bounds are checked by actually
# inflating, since the sizes a zip header declares can lie.
WORKBOOK_MAX_UNCOMPRESSED_BYTES = 150_000_000
WORKBOOK_MAX_COMPRESSION_RATIO = 100
WORKBOOK_MAX_ARCHIVE_MEMBERS = 10_000
MAX_WORKBOOK_DEFINED_NAMES = 500
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
# maxCells bounds the cell count but not the text: one long shared string
# referenced from thousands of cells is a small file and a huge response.
# Strings are cut per cell, and text (values, formulas, number formats,
# defined-name targets) is bounded per call.
MAX_WORKBOOK_CELL_CHARS = 2_000
MAX_WORKBOOK_TOTAL_CHARS = 200_000
WORKBOOK_TEXT_CUT_MARKER = "…"
MAX_WORKBOOK_RANGES = 20
# Measuring a sheet's used range reads every stored cell. A crafted sheet can
# hold millions, so the scan stops after this many cells or seconds and the
# extent is reported as incomplete.
WORKBOOK_EXTENT_MAX_CELLS = 5_000_000
WORKBOOK_EXTENT_MAX_SECONDS = 10.0
# Each range is read in up to three passes (values, formulas, number formats)
# from the top of the sheet, and the worker thread cannot be cancelled, so a
# call stops starting new passes after this long and marks the rest truncated.
WORKBOOK_READ_MAX_SECONDS = 30.0
# A defined name's target echoed in an error message is cut to this length.
MAX_WORKBOOK_ERROR_ECHO_CHARS = 200
WORKBOOK_ERROR_VALUE = re.compile(r"#(?:N/A|REF!|NAME\?|VALUE!|DIV/0!|NULL!|NUM!)")
# A bracketed workbook before a sheet reference ([1]Proforma!C4); structured
# table references like Table1[Rent] have no "!" and are kept.
WORKBOOK_EXTERNAL_REFERENCE = re.compile(r"\[[^\]]+\][^!\[]*!")

Bounds = tuple[int, int, int, int]


class WorkbookReadError(RuntimeError):
    """Raised when a spreadsheet or one of its ranges cannot be read; surfaced
    as ``unsupportedReason`` or a per-range ``error``."""


class _TextBudget:
    """Characters one call may still return, and when it must stop reading.

    Every string that reaches the response is charged: cell text, formulas,
    number formats, defined names and their targets, and sheet names. Numbers
    and booleans are charged their printed length; empty cells are bounded by
    maxCells instead."""

    def __init__(self, total: int) -> None:
        self.remaining = total
        self.cut_cells = 0
        self.deadline = monotonic() + WORKBOOK_READ_MAX_SECONDS

    def take(self, value: Any) -> Any:
        """Charge a cell value, cutting long text."""
        if value is None:
            return value
        if not isinstance(value, str):
            self.remaining -= len(str(value))
            return value
        if len(value) > MAX_WORKBOOK_CELL_CHARS:
            value = value[:MAX_WORKBOOK_CELL_CHARS] + WORKBOOK_TEXT_CUT_MARKER
            self.cut_cells += 1
        self.remaining -= len(value)
        return value

    def charge(self, text: str | None) -> None:
        """Charge an identifier (sheet or name) that must come back whole."""
        self.remaining -= len(text or "")

    @property
    def exhausted(self) -> bool:
        return self.remaining < 0

    @property
    def out_of_time(self) -> bool:
        return monotonic() > self.deadline


class WorkbookRead(NamedTuple):
    sheets: list[WorkbookSheetInfo]
    definedNameCount: int | None
    definedNamesSkipped: int | None
    definedNames: list[WorkbookDefinedNameInfo]
    ranges: list[AttachmentRangeData]
    truncated: bool


def unavailable_reason() -> str | None:
    if openpyxl is None:
        return "Workbook reading requires the openpyxl package"
    if not openpyxl.DEFUSEDXML:
        return "Workbook reading requires the defusedxml package"
    return None


def clamp_max_bytes(maxBytes: int) -> int:
    return max(1, min(maxBytes, MAX_WORKBOOK_MAX_BYTES))


def clamp_max_cells(maxCells: int) -> int:
    return max(1, min(maxCells, MAX_WORKBOOK_MAX_CELLS))


def is_spreadsheet(name: str | None, content_type: str | None) -> bool:
    content_type = (content_type or "").split(";")[0].strip().lower()
    return content_type in SPREADSHEET_CONTENT_TYPES or (name or "").lower().endswith(
        SPREADSHEET_EXTENSIONS
    )


def is_legacy_spreadsheet(name: str | None, content_type: str | None) -> bool:
    content_type = (content_type or "").split(";")[0].strip().lower()
    name = (name or "").lower()
    if name.endswith(SPREADSHEET_EXTENSIONS):
        return False
    return content_type in LEGACY_SPREADSHEET_CONTENT_TYPES or name.endswith(
        LEGACY_SPREADSHEET_EXTENSIONS
    )


def read_workbook(
    content_bytes: bytes,
    *,
    ranges: list[str] | None,
    sheet: str | None,
    includeFormulas: bool,
    includeNumberFormat: bool,
    includeLayout: bool,
    includeDefinedNames: bool,
    maxCells: int,
) -> WorkbookRead:
    reason = unavailable_reason()
    if reason is not None:
        raise WorkbookReadError(reason)
    _check_archive(content_bytes)
    books = _Books(content_bytes)
    try:
        return _read(
            books,
            ranges=ranges,
            sheet=sheet,
            includeFormulas=includeFormulas,
            includeNumberFormat=includeNumberFormat,
            includeLayout=includeLayout,
            includeDefinedNames=includeDefinedNames,
            maxCells=maxCells,
        )
    except WorkbookReadError:
        raise
    except Exception as error:  # pragma: no cover - parser internals
        raise WorkbookReadError(f"Could not read workbook: {error}") from error
    finally:
        books.close()


class _Books:
    """The same file opened twice, each only on first use: ``values`` holds
    Excel's cached results, ``formulas`` the formula text. A formula Excel
    never calculated is empty in ``values`` but not in ``formulas``, so the
    used range is measured on ``formulas``."""

    def __init__(self, content_bytes: bytes) -> None:
        self._content_bytes = content_bytes
        self._values: Any = None
        self._formulas: Any = None

    @property
    def values(self) -> Any:
        if self._values is None:
            self._values = _open(self._content_bytes, data_only=True)
        return self._values

    @property
    def formulas(self) -> Any:
        if self._formulas is None:
            self._formulas = _open(self._content_bytes, data_only=False)
        return self._formulas

    def close(self) -> None:
        for book in (self._values, self._formulas):
            if book is not None:
                book.close()


def _read(
    books: _Books,
    *,
    ranges: list[str] | None,
    sheet: str | None,
    includeFormulas: bool,
    includeNumberFormat: bool,
    includeLayout: bool,
    includeDefinedNames: bool,
    maxCells: int,
) -> WorkbookRead:
    # The <dimension> a writer records can be stale or swollen by formatting,
    # so the used range is measured from the non-empty cells. Measuring scans
    # the whole sheet, so it only happens for sheets the layout or a
    # whole-row/column request actually needs.
    extents: dict[str, Bounds | None] = {}
    incomplete_extents: set[str] = set()

    def extent_of(title: str) -> Bounds | None:
        if title not in extents:
            extents[title], complete = _worksheet_extent(books.formulas[title])
            if not complete:
                incomplete_extents.add(title)
        return extents[title]

    book = books.values
    sheets: list[WorkbookSheetInfo] = []
    defined_names: list[WorkbookDefinedNameInfo] = []
    name_count: int | None = None
    names_skipped: int | None = None
    truncated = False
    budget = _TextBudget(MAX_WORKBOOK_TOTAL_CHARS)
    if includeLayout or includeDefinedNames:
        usable_names, names_skipped = _workbook_defined_names(book)
        name_count = len(usable_names)
        if includeDefinedNames:
            truncated = name_count > MAX_WORKBOOK_DEFINED_NAMES
            for defined_name in usable_names[:MAX_WORKBOOK_DEFINED_NAMES]:
                budget.charge(defined_name.name)
                budget.charge(defined_name.scope)
                defined_name.value = budget.take(defined_name.value)
                if budget.exhausted:
                    truncated = True
                    break
                defined_names.append(defined_name)
    if includeLayout:
        for worksheet in book.worksheets:
            budget.charge(worksheet.title)
            if budget.exhausted:
                truncated = True
                break
            extent = extent_of(worksheet.title)
            sheets.append(
                WorkbookSheetInfo(
                    name=worksheet.title,
                    visibility=worksheet.sheet_state or "visible",
                    dimensions=_a1_address(*extent) if extent else None,
                    rowCount=extent[3] - extent[1] + 1 if extent else 0,
                    columnCount=extent[2] - extent[0] + 1 if extent else 0,
                    extentIncomplete=worksheet.title in incomplete_extents,
                )
            )

    requests: list[str | None] = list(ranges or [])
    if not requests and sheet:
        # A sheet without ranges means "its whole used range".
        requests = [None]

    range_data: list[AttachmentRangeData] = []
    remaining = maxCells
    for request in requests:
        try:
            data, used = _read_range(
                books,
                request,
                default_sheet=sheet,
                extent_of=extent_of,
                incomplete_extents=incomplete_extents,
                includeFormulas=includeFormulas,
                includeNumberFormat=includeNumberFormat,
                remaining=remaining,
                maxCells=maxCells,
                budget=budget,
            )
        except Exception as error:
            # One unreadable range never costs the caller the others.
            message = (
                str(error)
                if isinstance(error, WorkbookReadError)
                else f"Could not read '{request}': {error}"
            )
            data, used = (
                AttachmentRangeData(
                    worksheet=_requested_sheet(request) or sheet or "",
                    address=request or "",
                    error=message,
                ),
                0,
            )
        remaining -= used
        truncated = truncated or data.truncated
        range_data.append(data)

    return WorkbookRead(
        sheets=sheets,
        definedNameCount=name_count,
        definedNamesSkipped=names_skipped,
        definedNames=defined_names,
        ranges=range_data,
        truncated=truncated,
    )


def _read_range(
    books: _Books,
    request: str | None,
    *,
    default_sheet: str | None,
    extent_of: Callable[[str], Bounds | None],
    incomplete_extents: set[str],
    includeFormulas: bool,
    includeNumberFormat: bool,
    remaining: int,
    maxCells: int,
    budget: _TextBudget,
) -> tuple[AttachmentRangeData, int]:
    """Read one range within the remaining cell and text budgets; return it
    and the number of cells it used."""

    worksheet_name, bounds, used_extent = _resolve_range(
        books.values, request, default_sheet=default_sheet, extent_of=extent_of
    )
    # A whole-sheet/row/column request on a sheet whose scan stopped early
    # may be missing cells past the measured extent.
    extent_cut = used_extent and worksheet_name in incomplete_extents
    if bounds is None:
        return (
            AttachmentRangeData(
                worksheet=worksheet_name,
                address="",
                values=[],
                rowCount=0,
                columnCount=0,
                truncated=extent_cut,
            ),
            0,
        )

    def used_up(limit: str) -> tuple[AttachmentRangeData, int]:
        return (
            AttachmentRangeData(
                worksheet=worksheet_name,
                address=_a1_address(*bounds),
                truncated=True,
                error=f"{limit} is used up; request this range in another call",
            ),
            0,
        )

    budget.charge(worksheet_name)
    min_col, min_row, max_col, max_row = bounds
    column_count = max_col - min_col + 1
    row_count = max_row - min_row + 1
    range_truncated = extent_cut
    if row_count * column_count > remaining:
        allowed_rows = remaining // column_count
        if allowed_rows == 0:
            return used_up(f"maxCells={maxCells}")
        row_count = allowed_rows
        range_truncated = True
    if budget.remaining <= 0:
        return used_up(f"The {MAX_WORKBOOK_TOTAL_CHARS}-character text limit")
    if budget.out_of_time:
        return used_up(f"The {WORKBOOK_READ_MAX_SECONDS:g}-second read time limit")

    def grid(book: Any, read: Callable[[Any], Any], rows_wanted: int) -> list[list[Any]]:
        """Read up to ``rows_wanted`` rows, stopping before the row that
        exhausts the text budget."""
        rows: list[list[Any]] = []
        for row in book[worksheet_name].iter_rows(
            min_row=min_row,
            max_row=min_row + rows_wanted - 1,
            min_col=min_col,
            max_col=max_col,
        ):
            if len(rows) == rows_wanted:
                break
            cells = [budget.take(read(cell)) for cell in row][:column_count]
            if budget.exhausted or budget.out_of_time:
                return rows
            cells.extend([None] * (column_count - len(cells)))
            rows.append(cells)
        # Read-only worksheets stop at their last stored row, so pad to the
        # requested shape; rowCount and the address always match the values.
        rows.extend([None] * column_count for _ in range(rows_wanted - len(rows)))
        return rows

    def value(cell: Any) -> Any:
        return _cell_value(cell.value)

    cut_before = budget.cut_cells
    grids: list[list[list[Any]]] = []
    readers: list[tuple[Any, Callable[[Any], Any]]] = [(books.values, value)]
    if includeFormulas:
        readers.append((books.formulas, value))
    if includeNumberFormat:
        readers.append((books.values, lambda cell: getattr(cell, "number_format", None)))
    for book, read in readers:
        rows = grid(book, read, row_count) if not budget.out_of_time else []
        if len(rows) < row_count:
            # The text or time budget ran out; keep every grid the same height.
            row_count = len(rows)
            range_truncated = True
            grids = [earlier[:row_count] for earlier in grids]
        grids.append(rows)
        if row_count == 0:
            if budget.out_of_time:
                return used_up(f"The {WORKBOOK_READ_MAX_SECONDS:g}-second read time limit")
            return used_up(f"The {MAX_WORKBOOK_TOTAL_CHARS}-character text limit")
    values, *rest = grids
    formulas = rest.pop(0) if includeFormulas else None
    number_format = rest.pop(0) if includeNumberFormat else None
    # A cell whose text was cut also marks the range truncated.
    range_truncated = range_truncated or budget.cut_cells > cut_before

    return (
        AttachmentRangeData(
            worksheet=worksheet_name,
            address=_a1_address(min_col, min_row, max_col, min_row + row_count - 1),
            values=values,
            formulas=formulas,
            numberFormat=number_format,
            rowCount=row_count,
            columnCount=column_count,
            truncated=range_truncated,
        ),
        row_count * column_count,
    )


def _check_archive(content_bytes: bytes) -> None:
    try:
        check_zip_archive(
            content_bytes,
            max_inflated_bytes=WORKBOOK_MAX_UNCOMPRESSED_BYTES,
            max_ratio=WORKBOOK_MAX_COMPRESSION_RATIO,
            max_members=WORKBOOK_MAX_ARCHIVE_MEMBERS,
            label="Workbook",
        )
    except NotAZipError as error:
        raise WorkbookReadError(
            "Attachment is not a readable .xlsx workbook; it may be "
            "password-protected or a legacy .xls file with an .xlsx name"
        ) from error
    except ZipGuardError as error:
        raise WorkbookReadError(str(error)) from error


def _open(content_bytes: bytes, *, data_only: bool) -> Any:
    try:
        return openpyxl.load_workbook(
            io.BytesIO(content_bytes),
            read_only=True,
            data_only=data_only,
            keep_links=False,
        )
    except Exception as error:
        raise WorkbookReadError(f"Could not open workbook: {error}") from error


def _worksheet_extent(worksheet: Any) -> tuple[Bounds | None, bool]:
    """Return ``(min_col, min_row, max_col, max_row)`` of non-empty cells and
    whether the whole sheet was scanned (see WORKBOOK_EXTENT_MAX_*)."""

    worksheet.reset_dimensions()
    min_col = min_row = max_col = max_row = None
    deadline = monotonic() + WORKBOOK_EXTENT_MAX_SECONDS
    scanned = 0
    complete = True
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, values_only=True), start=1
    ):
        scanned += len(row)
        if scanned > WORKBOOK_EXTENT_MAX_CELLS or monotonic() > deadline:
            complete = False
            break
        for column_index, value in enumerate(row, start=1):
            if value is None or value == "":
                continue
            if min_row is None:
                min_row = row_index
            max_row = row_index
            if min_col is None or column_index < min_col:
                min_col = column_index
            if max_col is None or column_index > max_col:
                max_col = column_index
    if min_row is None:
        return None, complete
    return (min_col, min_row, max_col, max_row), complete


def _workbook_defined_names(book: Any) -> tuple[list[WorkbookDefinedNameInfo], int]:
    """Return the usable defined names and how many leftovers were skipped."""

    names: list[WorkbookDefinedNameInfo] = []
    skipped = 0

    def collect(defined_names: Any, scope: str | None) -> None:
        nonlocal skipped
        for defined_name in defined_names.values():
            if not _is_useful_defined_name(defined_name):
                skipped += 1
                continue
            names.append(
                WorkbookDefinedNameInfo(
                    name=defined_name.name,
                    value=defined_name.value or "",
                    scope=scope,
                )
            )

    collect(book.defined_names, None)
    for worksheet in book.worksheets:
        collect(worksheet.defined_names, worksheet.title)
    return names, skipped


def _is_useful_defined_name(defined_name: Any) -> bool:
    """Sponsor models picked up from broker templates carry hundreds of names
    that are noise when looking for a model's inputs: built-ins (print areas,
    filters), hidden names, underscore-prefixed print-macro leftovers, array
    constants, names whose target is an error (#N/A, deleted cells), and links
    into other workbooks this tool cannot read. They still resolve when
    requested by name; they are only unlisted."""

    name = defined_name.name or ""
    value = (defined_name.value or "").strip()
    if defined_name.is_reserved or getattr(defined_name, "hidden", False):
        return False
    if not value or name.startswith("_"):
        return False
    if value.startswith("{") or WORKBOOK_ERROR_VALUE.search(value):
        return False
    return not WORKBOOK_EXTERNAL_REFERENCE.search(value)


def _resolve_range(
    book: Any,
    request: str | None,
    *,
    default_sheet: str | None,
    extent_of: Callable[[str], Bounds | None],
) -> tuple[str, Bounds | None, bool]:
    """Turn an A1 address, ``Sheet!A1:B2``, a defined name, or ``Sheet!Name``
    into a sheet and ``(min_col, min_row, max_col, max_row)`` bounds. Whole
    rows or columns are clamped to the sheet's used range; ``None`` bounds
    mean the request covers no used cells. The flag is true when the bounds
    came from the measured used range."""

    sheet_name: str | None = None
    address: str | None = None
    if request is not None:
        text = request.strip()
        if "!" in text:
            sheet_part, _, address = text.rpartition("!")
            sheet_name = _unquote_sheet_name(sheet_part)
        else:
            address = text
        address = address.replace("$", "").strip()
        if address and not _is_a1_reference(address):
            # 'Sheet'!Name means a name scoped to that sheet, as in Excel.
            sheet_name, address = _resolve_defined_name(
                book, address, scope_sheet=sheet_name or default_sheet
            )

    worksheet_name = _match_worksheet(book, sheet_name or default_sheet)
    if not address:
        return worksheet_name, extent_of(worksheet_name), True

    try:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except (TypeError, ValueError) as error:
        raise WorkbookReadError(
            f"'{request}' is not an A1 address or defined name"
        ) from error
    # openpyxl parses A0 as row 0; iter_rows would read row 1 under that label.
    if any(bound is not None and bound < 1 for bound in (min_col, min_row, max_col, max_row)):
        raise WorkbookReadError(
            f"'{request}' is not a valid address; rows and columns start at 1"
        )
    used_extent = None in (min_col, min_row, max_col, max_row)
    if used_extent:
        extent = extent_of(worksheet_name)
        if extent is None:
            return worksheet_name, None, True
        min_col = min_col or extent[0]
        min_row = min_row or extent[1]
        max_col = max_col or extent[2]
        max_row = max_row or extent[3]
    # Excel reads A10001:A1 as A1:A10001. Normalising here also keeps every
    # cell count positive, so a reversed range can never add to the budget.
    min_col, max_col = sorted((min_col, max_col))
    min_row, max_row = sorted((min_row, max_row))
    if max_row > EXCEL_MAX_ROWS or max_col > EXCEL_MAX_COLUMNS:
        raise WorkbookReadError(
            f"'{request}' is outside a worksheet ({EXCEL_MAX_ROWS} rows, "
            f"{EXCEL_MAX_COLUMNS} columns)"
        )
    return worksheet_name, (min_col, min_row, max_col, max_row), used_extent


def _match_worksheet(book: Any, sheet_name: str | None) -> str:
    titles = book.sheetnames
    if sheet_name is None:
        if len(titles) == 1:
            return titles[0]
        raise WorkbookReadError(
            "Workbook has several sheets; qualify the address with a sheet "
            "name (e.g. 'Unit Mix'!A1:H40) or pass sheet"
        )
    if sheet_name in titles:
        return sheet_name
    # Excel treats sheet names case-insensitively.
    for title in titles:
        if title.casefold() == sheet_name.casefold():
            return title
    raise WorkbookReadError(f"Workbook has no sheet named '{sheet_name}'")


def _resolve_defined_name(
    book: Any,
    name: str,
    *,
    scope_sheet: str | None,
) -> tuple[str, str]:
    candidates = list(book.defined_names.items())
    if scope_sheet is not None:
        try:
            scoped = book[_match_worksheet(book, scope_sheet)]
        except WorkbookReadError:
            scoped = None
        if scoped is not None:
            # A sheet-scoped name shadows a workbook name of the same name.
            candidates = list(scoped.defined_names.items()) + candidates
    for candidate_name, defined_name in candidates:
        if candidate_name.casefold() != name.casefold():
            continue
        try:
            destinations = list(defined_name.destinations)
        except Exception:
            # openpyxl cannot parse some targets, e.g. structured table
            # references like Rents[Rent].
            destinations = []
        target = _echo(defined_name.value)
        if len(destinations) != 1:
            raise WorkbookReadError(
                f"Defined name '{name}' refers to '{target}', which "
                "is not a single cell range this tool can read"
            )
        sheet_name, address = destinations[0]
        address = address.replace("$", "").strip()
        # openpyxl reads Sheet!#REF! as that sheet with an empty address, which
        # would otherwise mean "the whole used range".
        if not address or "#" in address:
            raise WorkbookReadError(
                f"Defined name '{name}' refers to '{target}', whose cells were "
                "deleted (#REF!)"
            )
        return sheet_name.replace("''", "'"), address
    raise WorkbookReadError(f"'{name}' is not an A1 address or defined name")


def _echo(text: str | None) -> str:
    """Cut workbook text quoted in an error message."""
    text = text or ""
    if len(text) <= MAX_WORKBOOK_ERROR_ECHO_CHARS:
        return text
    return text[:MAX_WORKBOOK_ERROR_ECHO_CHARS] + WORKBOOK_TEXT_CUT_MARKER


def _requested_sheet(request: str | None) -> str | None:
    """The sheet named in a ``Sheet!A1`` request, for labelling its error."""
    if request is None or "!" not in request:
        return None
    return _unquote_sheet_name(request.strip().rpartition("!")[0]) or None


def _unquote_sheet_name(sheet_part: str) -> str:
    sheet_part = sheet_part.strip()
    if len(sheet_part) >= 2 and sheet_part[0] == sheet_part[-1] == "'":
        return sheet_part[1:-1].replace("''", "'")
    return sheet_part


def _is_a1_reference(address: str) -> bool:
    try:
        bounds = range_boundaries(address)
    except (TypeError, ValueError):
        return False
    # Without a colon only a full cell (B5) is an address; a bare "IRR" is a
    # defined name in Excel, not column IRR.
    return ":" in address or None not in bounds


def _a1_address(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    # Array and data-table formulas come back as objects carrying the text.
    text = getattr(value, "text", None)
    return text if isinstance(text, str) else str(value)
